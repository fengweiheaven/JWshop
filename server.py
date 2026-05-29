from __future__ import annotations

import argparse
import csv
import json
import math
import re
import sys
import threading
import webbrowser
from datetime import date, datetime, time
from email import policy
from email.parser import BytesParser
from http import HTTPStatus
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
    from openpyxl.utils.datetime import from_excel
except ImportError:  # pragma: no cover - surfaced through the API response.
    load_workbook = None
    from_excel = None


def get_resource_dir() -> Path:
    if getattr(sys, "frozen", False) and hasattr(sys, "_MEIPASS"):
        return Path(sys._MEIPASS)
    return Path(__file__).resolve().parent


def get_runtime_dir() -> Path:
    if getattr(sys, "frozen", False):
        if sys.platform == "darwin":
            path = Path.home() / "Library" / "Application Support" / "InventoryTool"
            path.mkdir(parents=True, exist_ok=True)
            return path
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


BASE_DIR = get_resource_dir()
RUNTIME_DIR = get_runtime_dir()
DEFAULT_PORT = 8765

COLUMN_ALIASES = {
    "date": [
        "日期",
        "时间",
        "业务日期",
        "单据日期",
        "库存日期",
        "统计日期",
        "入库日期",
        "出库日期",
        "订单日期",
        "创建日期",
    ],
    "product": [
        "商品名称",
        "商品名",
        "商品",
        "产品名称",
        "产品名",
        "品名",
        "货品名称",
        "物料名称",
        "名称",
    ],
    "case_qty": [
        "整件数量",
        "整件",
        "整箱数量",
        "整箱",
        "库存件数",
        "件数",
        "箱数",
        "件",
        "箱",
    ],
    "bottle_qty": [
        "单瓶数量",
        "单瓶",
        "散瓶数量",
        "散瓶",
        "零瓶数量",
        "零瓶",
        "瓶数",
        "支数",
        "数量(瓶)",
        "数量（瓶）",
        "瓶",
    ],
    "stock_qty": [
        "剩余库存",
        "库存数量",
        "当前库存",
        "库存",
        "结存",
        "期末库存",
        "可用库存",
    ],
    "box_size": [
        "箱规",
        "包装规格",
        "规格",
        "件装",
        "箱装",
        "装箱数",
        "每箱数量",
        "每箱瓶数",
        "瓶/箱",
        "瓶每箱",
    ],
}

COLUMN_LABELS = {
    "date": "日期",
    "product": "商品名称",
    "case_qty": "整件数量",
    "bottle_qty": "单瓶数量",
    "stock_qty": "库存数量",
    "box_size": "箱规",
}


class UserFacingError(Exception):
    pass


def normalize_label(value: Any) -> str:
    text = "" if value is None else str(value)
    return re.sub(r"[\s_\-:：/\\（）()\[\]【】,，.。]+", "", text).lower()


def is_empty_row(row: list[Any] | tuple[Any, ...]) -> bool:
    return all(cell is None or str(cell).strip() == "" for cell in row)


def alias_score(header: Any, aliases: list[str]) -> int:
    label = normalize_label(header)
    if not label:
        return 0

    best = 0
    for index, alias in enumerate(aliases):
        normalized_alias = normalize_label(alias)
        if label == normalized_alias:
            best = max(best, 1000 - index)
        elif normalized_alias and normalized_alias in label:
            best = max(best, 500 - index)
    return best


def detect_columns(headers: list[Any]) -> tuple[dict[str, int | None], dict[str, int]]:
    detected: dict[str, int | None] = {}
    scores: dict[str, int] = {}
    used_indexes: set[int] = set()

    for field, aliases in COLUMN_ALIASES.items():
        candidates: list[tuple[int, int]] = []
        for index, header in enumerate(headers):
            if index in used_indexes:
                continue
            score = alias_score(header, aliases)
            if score:
                candidates.append((score, index))

        if candidates:
            score, index = max(candidates, key=lambda item: (item[0], -item[1]))
            detected[field] = index
            scores[field] = score
            used_indexes.add(index)
        else:
            detected[field] = None
            scores[field] = 0

    return detected, scores


def find_header_row(rows: list[list[Any]]) -> tuple[int, list[Any], dict[str, int | None]]:
    candidates = rows[: min(30, len(rows))]
    best_index = -1
    best_headers: list[Any] = []
    best_detected: dict[str, int | None] = {}
    best_rank = (-1, -1)

    for index, row in enumerate(candidates):
        if is_empty_row(row):
            continue
        detected, scores = detect_columns(list(row))
        hits = sum(1 for column in detected.values() if column is not None)
        score_total = sum(scores.values())
        rank = (hits, score_total)
        if rank > best_rank:
            best_index = index
            best_headers = list(row)
            best_detected = detected
            best_rank = rank

    if best_index < 0 or best_rank[0] < 2:
        raise UserFacingError("没有识别到表头。请确认表格包含日期、商品名称、整件数量、单瓶数量、箱规等列。")

    return best_index, best_headers, best_detected


def parse_date_value(value: Any) -> str | None:
    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return None

    if isinstance(value, (int, float)) and from_excel:
        try:
            parsed = from_excel(value)
            if isinstance(parsed, datetime):
                return parsed.date().isoformat()
        except Exception:
            return None

    text = str(value).strip()
    if not text:
        return None

    text = re.sub(r"\s+", " ", text)
    text = text.replace("年", "-").replace("月", "-").replace("日", "")
    text = text.replace(".", "-").replace("/", "-")
    match = re.search(r"(\d{4}-\d{1,2}-\d{1,2})", text)
    if match:
        text = match.group(1)

    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
        try:
            return datetime.strptime(text[: len(fmt)], fmt).date().isoformat()
        except ValueError:
            continue

    # Handles compact strings like 20260528.
    compact = re.fullmatch(r"(\d{4})(\d{2})(\d{2})", re.sub(r"\D", "", text))
    if compact:
        year, month, day = compact.groups()
        try:
            return date(int(year), int(month), int(day)).isoformat()
        except ValueError:
            return None

    return None


def parse_number(value: Any, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        if math.isnan(float(value)):
            return default
        return float(value)

    text = str(value).strip().replace(",", "")
    if not text:
        return default

    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return default
    try:
        return float(match.group(0))
    except ValueError:
        return default


def parse_box_size(value: Any) -> float:
    if value is None or value == "":
        return 1.0
    if isinstance(value, (int, float)) and float(value) > 0:
        return float(value)

    text = str(value).strip().replace(",", "")
    numbers = re.findall(r"\d+(?:\.\d+)?", text)
    if not numbers:
        return 1.0

    # Specifications often look like "500ml*24瓶"; the last positive number is
    # usually the number of bottles per case.
    for raw_number in reversed(numbers):
        number = float(raw_number)
        if number > 0:
            return number
    return 1.0


def cell(row: list[Any], index: int | None) -> Any:
    if index is None or index >= len(row):
        return None
    return row[index]


def normalize_product_unit(product: str) -> tuple[str, str]:
    text = product.strip()
    bottle_patterns = (
        r"\s*[/／]\s*瓶\s*$",
        r"\s*[（(]\s*瓶\s*[）)]\s*$",
        r"\s*[-_]\s*瓶\s*$",
    )

    for pattern in bottle_patterns:
        if re.search(pattern, text):
            return re.sub(pattern, "", text).strip(), "bottle"
    return text, "case"


def detect_movement_columns(
    headers: list[Any],
    detected: dict[str, int | None],
) -> list[dict[str, Any]]:
    skipped_indexes = {index for index in detected.values() if index is not None}
    movement_columns: list[dict[str, Any]] = []

    customer_return_keywords = ("客退", "退货", "退回", "退库", "退仓", "退件")
    surplus_return_keywords = ("多货实退", "多货退货", "多货退", "实退")
    outbound_keywords = (
        "出库",
        "补货",
        "夜补",
        "正转残",
        "发货",
        "送货",
        "配送",
        "销售",
        "售出",
        "调货",
    )
    excluded_keywords = ("总数", "合计", "小计", "费用", "费", "库存", "备注")

    for index, header in enumerate(headers):
        if index in skipped_indexes:
            continue

        label = "" if header is None else str(header).strip()
        normalized = normalize_label(label)
        if not normalized or any(keyword in normalized for keyword in excluded_keywords):
            continue

        if any(keyword in normalized for keyword in surplus_return_keywords):
            movement_columns.append(
                {"index": index, "label": label, "kind": "surplus_return", "sign": 1}
            )
        elif any(keyword in normalized for keyword in customer_return_keywords):
            movement_columns.append(
                {"index": index, "label": label, "kind": "customer_return", "sign": 1}
            )
        elif any(keyword in normalized for keyword in outbound_keywords):
            movement_columns.append({"index": index, "label": label, "kind": "outbound", "sign": -1})

    return movement_columns


def convert_movement_quantity(value: Any, unit_type: str, box_size: float) -> float:
    quantity = abs(parse_number(value))
    if not quantity:
        return 0.0
    if unit_type == "bottle":
        return quantity / box_size if box_size else 0.0
    return quantity


def read_xlsx_rows(content: bytes) -> tuple[str, list[str], list[list[Any]]]:
    if load_workbook is None:
        raise UserFacingError("当前 Python 环境缺少 openpyxl，无法读取 Excel。请先安装 openpyxl。")

    workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    sheet_names = workbook.sheetnames

    selected_title = ""
    selected_rows: list[list[Any]] = []
    for worksheet in workbook.worksheets:
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        meaningful_count = sum(1 for row in rows if not is_empty_row(row))
        if meaningful_count > 1:
            selected_title = worksheet.title
            selected_rows = rows
            break

    if not selected_rows:
        raise UserFacingError("这个 Excel 里没有可读取的数据行。")

    return selected_title, sheet_names, selected_rows


def read_csv_rows(content: bytes) -> tuple[str, list[str], list[list[Any]]]:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError as exc:
            last_error = exc
    else:
        raise UserFacingError(f"CSV 编码无法识别：{last_error}")

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample)
    except csv.Error:
        dialect = csv.excel

    reader = csv.reader(StringIO(text), dialect)
    rows = [list(row) for row in reader]
    if len(rows) < 2:
        raise UserFacingError("CSV 里没有足够的数据行。")
    return "CSV", ["CSV"], rows


def parse_table(filename: str, content: bytes) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        sheet_name, sheet_names, rows = read_xlsx_rows(content)
    elif suffix == ".csv":
        sheet_name, sheet_names, rows = read_csv_rows(content)
    elif suffix == ".xls":
        raise UserFacingError("暂不支持旧版 .xls，请另存为 .xlsx 或 .csv 后再上传。")
    else:
        raise UserFacingError("请上传 .xlsx、.xlsm 或 .csv 文件。")

    header_index, headers, detected = find_header_row(rows)
    movement_columns = detect_movement_columns(headers, detected)
    data_rows = rows[header_index + 1 :]

    has_direct_quantity = detected["case_qty"] is not None or detected["bottle_qty"] is not None
    has_split_stock = detected.get("stock_qty") is not None
    quantity_mode = "direct" if has_direct_quantity else "split_stock"

    if not has_direct_quantity and not has_split_stock:
        raise UserFacingError("没有识别到库存数量列。请确认表格包含整件/单瓶数量列，或包含“剩余库存/库存”列。")

    required_fields = ["date", "product", "box_size"]
    if quantity_mode == "direct":
        required_fields.extend(["case_qty", "bottle_qty"])
    else:
        required_fields.append("stock_qty")
    missing = [COLUMN_LABELS[field] for field in required_fields if detected.get(field) is None]

    normalized_rows: list[dict[str, Any]] = []
    skipped_rows = 0
    current_date: str | None = None

    for offset, row in enumerate(data_rows, start=header_index + 2):
        if is_empty_row(row):
            continue

        parsed_date = parse_date_value(cell(row, detected["date"]))
        if parsed_date:
            current_date = parsed_date

        row_date = parsed_date or current_date
        raw_product = str(cell(row, detected["product"]) or "").strip()
        if not row_date or not raw_product:
            skipped_rows += 1
            continue

        product, unit_type = normalize_product_unit(raw_product)
        box_size = parse_box_size(cell(row, detected["box_size"]))

        if quantity_mode == "split_stock":
            stock_qty = parse_number(cell(row, detected["stock_qty"]))
            case_qty = stock_qty if unit_type == "case" else 0.0
            bottle_qty = stock_qty if unit_type == "bottle" else 0.0
        else:
            case_qty = parse_number(cell(row, detected["case_qty"]))
            bottle_qty = parse_number(cell(row, detected["bottle_qty"]))

        total_qty = case_qty + (bottle_qty / box_size if box_size else 0)
        outbound_qty = 0.0
        return_qty = 0.0
        outbound_case_qty = 0.0
        outbound_bottle_qty = 0.0
        customer_return_qty = 0.0
        customer_return_case_qty = 0.0
        customer_return_bottle_qty = 0.0
        surplus_return_qty = 0.0
        surplus_return_case_qty = 0.0
        surplus_return_bottle_qty = 0.0

        for movement_column in movement_columns:
            raw_movement_qty = abs(parse_number(cell(row, movement_column["index"])))
            movement_qty = convert_movement_quantity(raw_movement_qty, unit_type, box_size)

            if movement_column["kind"] == "customer_return":
                customer_return_qty += movement_qty
                return_qty += movement_qty
                if unit_type == "bottle":
                    customer_return_bottle_qty += raw_movement_qty
                else:
                    customer_return_case_qty += raw_movement_qty
            elif movement_column["kind"] == "surplus_return":
                surplus_return_qty += movement_qty
                return_qty += movement_qty
                if unit_type == "bottle":
                    surplus_return_bottle_qty += raw_movement_qty
                else:
                    surplus_return_case_qty += raw_movement_qty
            else:
                outbound_qty += movement_qty
                if unit_type == "bottle":
                    outbound_bottle_qty += raw_movement_qty
                else:
                    outbound_case_qty += raw_movement_qty

        normalized_rows.append(
            {
                "sourceRow": offset,
                "date": row_date,
                "product": product,
                "caseQty": case_qty,
                "bottleQty": bottle_qty,
                "boxSize": box_size,
                "totalQty": total_qty,
                "outboundQty": outbound_qty,
                "outboundCaseQty": outbound_case_qty,
                "outboundBottleQty": outbound_bottle_qty,
                "returnQty": return_qty,
                "customerReturnQty": customer_return_qty,
                "customerReturnCaseQty": customer_return_case_qty,
                "customerReturnBottleQty": customer_return_bottle_qty,
                "surplusReturnQty": surplus_return_qty,
                "surplusReturnCaseQty": surplus_return_case_qty,
                "surplusReturnBottleQty": surplus_return_bottle_qty,
                "changeQty": return_qty - outbound_qty,
            }
        )

    if not normalized_rows:
        raise UserFacingError("没有可统计的数据。请检查日期和商品名称列是否有内容。")

    detected_columns = {
        field: (str(headers[index]).strip() if index is not None and index < len(headers) else None)
        for field, index in detected.items()
    }
    dates = sorted({row["date"] for row in normalized_rows})

    return {
        "filename": filename,
        "sheetName": sheet_name,
        "sheetNames": sheet_names,
        "headerRow": header_index + 1,
        "detectedColumns": detected_columns,
        "quantityMode": quantity_mode,
        "missingColumns": missing,
        "movementColumns": [
            {"label": column["label"], "kind": column["kind"]} for column in movement_columns
        ],
        "rowCount": len(normalized_rows),
        "skippedRows": skipped_rows,
        "dates": dates,
        "rows": normalized_rows,
    }


def parse_multipart_file(headers: Any, body: bytes) -> tuple[str, bytes]:
    content_type = headers.get("Content-Type", "")
    if "multipart/form-data" not in content_type:
        raise UserFacingError("请求格式不正确，请重新上传文件。")

    raw_message = (
        f"Content-Type: {content_type}\r\nMIME-Version: 1.0\r\n\r\n".encode("utf-8")
        + body
    )
    message = BytesParser(policy=policy.default).parsebytes(raw_message)

    for part in message.iter_parts():
        name = part.get_param("name", header="content-disposition")
        filename = part.get_param("filename", header="content-disposition")
        if name == "file" and filename:
            content = part.get_payload(decode=True)
            if not content:
                raise UserFacingError("上传的文件是空的。")
            return filename, content

    raise UserFacingError("没有收到表格文件。")


class InventoryToolHandler(SimpleHTTPRequestHandler):
    server_version = "InventoryTool/1.0"

    def __init__(self, *args: Any, **kwargs: Any) -> None:
        super().__init__(*args, directory=str(BASE_DIR), **kwargs)

    def do_POST(self) -> None:
        if self.path != "/api/analyze":
            self.send_error(HTTPStatus.NOT_FOUND)
            return

        try:
            content_length = int(self.headers.get("Content-Length", "0"))
            body = self.rfile.read(content_length)
            filename, content = parse_multipart_file(self.headers, body)
            payload = parse_table(filename, content)
            self.send_json({"ok": True, "data": payload})
        except UserFacingError as exc:
            self.send_json({"ok": False, "error": str(exc)}, HTTPStatus.BAD_REQUEST)
        except Exception as exc:  # pragma: no cover - keeps UI friendly.
            self.send_json({"ok": False, "error": f"处理失败：{exc}"}, HTTPStatus.INTERNAL_SERVER_ERROR)

    def end_headers(self) -> None:
        self.send_header("Cache-Control", "no-store")
        super().end_headers()

    def send_json(self, payload: dict[str, Any], status: HTTPStatus = HTTPStatus.OK) -> None:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)


def ensure_background_logs() -> None:
    if sys.stdout is None:
        sys.stdout = (RUNTIME_DIR / "server.out.log").open("a", encoding="utf-8", buffering=1)
    if sys.stderr is None:
        sys.stderr = (RUNTIME_DIR / "server.err.log").open("a", encoding="utf-8", buffering=1)


def open_browser_later(port: int) -> None:
    url = f"http://127.0.0.1:{port}/"
    threading.Timer(0.8, lambda: webbrowser.open(url)).start()


def create_server(preferred_port: int) -> tuple[ThreadingHTTPServer, int]:
    ports = [preferred_port]
    if preferred_port == DEFAULT_PORT:
        ports.extend(range(DEFAULT_PORT + 1, DEFAULT_PORT + 20))

    last_error: OSError | None = None
    for port in ports:
        try:
            return ThreadingHTTPServer(("127.0.0.1", port), InventoryToolHandler), port
        except OSError as exc:
            last_error = exc

    raise last_error or OSError("没有可用端口")


def main() -> None:
    ensure_background_logs()
    parser = argparse.ArgumentParser(description="Local inventory summary tool")
    parser.add_argument("--port", type=int, default=DEFAULT_PORT)
    parser.add_argument("--no-browser", action="store_true")
    args = parser.parse_args()

    server, port = create_server(args.port)
    if not args.no_browser:
        open_browser_later(port)
    if sys.stdout:
        print(f"库存汇总工具已启动：http://127.0.0.1:{port}")
        print("按 Ctrl+C 停止服务。")
    server.serve_forever()


if __name__ == "__main__":
    main()
